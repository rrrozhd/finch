"""Runs INSIDE the Daytona sandbox: unzip untrusted upload, parse csv/xlsx, OCR receipt images with confidence.

Usage: python sandbox_worker.py upload.zip outdir  -> prints one JSON document to stdout.
No API keys live here; low-confidence receipts are flagged for the vision fallback done by the server.
"""
import csv
import io
import json
import os
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

MAX_FILES, MAX_BYTES = 200, 50_000_000
CONF_GATE = 0.80  # weakest OCR line score below this -> vision fallback
MONEY = re.compile(r"\$?\s*([0-9]{1,3}(?:,[0-9]{3})*\.[0-9]{2})")


def safe_extract(zpath: Path, out: Path) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    files, total = [], 0
    with zipfile.ZipFile(zpath) as z:
        for info in z.infolist()[:MAX_FILES]:
            name = Path(info.filename).name  # drop directories: no zip-slip
            if info.is_dir() or not name or name.startswith("."):
                continue
            total += info.file_size
            if total > MAX_BYTES:
                raise SystemExit("upload too large")
            p = out / name
            p.write_bytes(z.read(info))
            files.append(p)
    return files


def parse_csv(p: Path) -> list[dict]:
    rows = [{(k or "").strip().lower(): v for k, v in r.items()} for r in csv.DictReader(io.StringIO(p.read_text(errors="replace")))]
    for r in rows:  # bank export dialects: Chase uses "Transaction Date"/"Posting Date" and MM/DD/YYYY
        r.setdefault("date", r.get("transaction date") or r.get("posting date"))
        if re.fullmatch(r"\d{2}/\d{2}/\d{4}", r.get("date") or ""):
            r["date"] = datetime.strptime(r["date"], "%m/%d/%Y").date().isoformat()
    need = {"date", "description", "amount"}
    if not rows or not need <= set(rows[0]) or "_answer_key" in p.parts:
        return []
    return [dict(id=r.get("transaction_id") or f"{p.name}#{i}", date=r["date"], description=r["description"].strip(),
                 amount=float(r["amount"]), source=f"{p.name}#{r.get('transaction_id') or i}") for i, r in enumerate(rows, 1)]


def parse_xlsx(p: Path) -> list[dict]:
    from openpyxl import load_workbook
    ws = load_workbook(p, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "").lower() for h in rows[0]]
    return [dict(zip(head, r), source=f"{p.name}!row{i}") for i, r in enumerate(rows[1:], 2) if any(r)]


_ocr = None


def ocr(p: Path) -> dict:
    global _ocr
    if _ocr is None:
        from rapidocr_onnxruntime import RapidOCR
        # the sandbox reports the host's core count; uncapped ONNX threads thrash a small CPU quota (69s vs 3s per image)
        n = int(os.getenv("OCR_THREADS", "2"))
        _ocr = RapidOCR(intra_op_num_threads=n, inter_op_num_threads=1)
    result, _ = _ocr(str(p))
    lines = [(t, float(s)) for _, t, s in (result or [])]
    text = "\n".join(t for t, _ in lines)
    conf = min((s for _, s in lines), default=0.0)  # weakest line: one smudged total matters more than the average
    total = None
    for i, (t, _) in enumerate(lines):  # amount on the "Total" line, or the next line
        if t.strip().lower().startswith("total"):
            m = MONEY.search(t) or (MONEY.search(lines[i + 1][0]) if i + 1 < len(lines) else None)
            total = float(m.group(1).replace(",", "")) if m else None
    d = None
    m = re.search(r"Date:?\s*([A-Z][a-z]{2})\s*(\d{1,2}),?\s*(\d{4})", text)  # OCR often drops spaces
    if m:
        d = datetime.strptime(" ".join(m.groups()), "%b %d %Y").date().isoformat()
    vendor = next((lines[i + 1][0] for i, (t, _) in enumerate(lines) if t.strip() == "From" and i + 1 < len(lines)), None)
    needs_vision = conf < CONF_GATE or total is None or d is None
    return dict(file=p.name, confidence=round(conf, 3), total=total, date=d, vendor=vendor,
                needs_vision=needs_vision, text=text[:2000], source=p.name)


def main(zpath: str, outdir: str):
    files = safe_extract(Path(zpath), Path(outdir))
    out = dict(transactions=[], sheets=[], receipts=[], skipped=[])
    for p in files:
        ext = p.suffix.lower()
        if ext == ".csv":
            out["transactions"] += parse_csv(p)
        elif ext in (".xlsx", ".xlsm"):
            out["sheets"].append(dict(file=p.name, rows=parse_xlsx(p)))
        elif ext in (".png", ".jpg", ".jpeg"):
            out["receipts"].append(ocr(p))
        else:
            out["skipped"].append(p.name)
    print(json.dumps(out, default=str))


if __name__ == "__main__":
    main(*sys.argv[1:3])
